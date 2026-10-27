"""Shared paths and read-only xlsx helpers for the Claude framework scripts."""
import sys
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "_docs"
TEMPLATES = ROOT / ".claude" / "templates"
REQS_XLSX = DOCS / "0.requirements-analysis.xlsx"

CATALOG_SHEET = "Capacity Catalog"
CATALOG_ID = "ID"
CATALOG_NAME = "Capacity Full Name"
REQ_ID = "Req ID"
REQ_CAPACITY = "Capacity Name"
STATUSES = ("0. Idea", "1. Draft", "2. Approve", "3. Implement", "4. Done")

DECISIONS_SHEET = "Open Decisions"
DECISION_ID = "Decision ID"
DECISION_STATUSES = ("Open", "Decided", "Deferred")

sys.stdout.reconfigure(encoding="utf-8")


def fail(message):
    print(f"error: {message}", file=sys.stderr)
    sys.exit(1)


def load_workbook(path=REQS_XLSX):
    import openpyxl

    with warnings.catch_warnings():
        # openpyxl warns that it does not support the Excel list validation; reading is not affected.
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path)


def headers(ws):
    return [cell.value for cell in ws[1]]


def read_rows(ws):
    """Return non-empty rows as dicts keyed by header, plus the row number in `_row`."""
    keys = headers(ws)
    rows = []
    for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not values[0]:
            continue
        row = {key: value for key, value in zip(keys, values) if key}
        row["_row"] = number
        rows.append(row)
    return rows


def theme_sheets(wb):
    return [ws for ws in wb.worksheets if ws.title not in (CATALOG_SHEET, DECISIONS_SHEET)]
